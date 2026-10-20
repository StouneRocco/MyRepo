import datetime as dt
import importlib
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from import_temperatures import RESOURCE_DIR, measurement, read_rows, insert_measurements


class ImportTests(unittest.TestCase):
    def test_bulk_import_preserves_duplicates_and_reports_other_errors(self):
        from pymongo.errors import BulkWriteError
        collection = MagicMock()
        collection.insert_many.side_effect = BulkWriteError({'nInserted':1,'writeErrors':[{'code':11000}]})
        self.assertEqual(insert_measurements(collection,[{'room':'a'},{'room':'b'}]),1)
        collection.insert_many.side_effect = BulkWriteError({'nInserted':0,'writeErrors':[{'code':121}]})
        with self.assertRaises(BulkWriteError):
            insert_measurements(collection,[{'room':'a'}])
    def test_new_monthly_semicolon_csv(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'month.csv'
            path.write_text('timestamp;display name;sensor id;value;room;floor;zone\n2026-06-30 23:55:00.021;Temperature;sensor;24.5;03-23-N;Etage 3;ESEO\n')
            doc = measurement(next(read_rows(path)))
            self.assertEqual(doc['temperature'], 24.5)
            self.assertEqual(doc['timestamp'], dt.datetime(2026,6,30,21,55,0,21000,tzinfo=dt.timezone.utc))

    def test_projection_context_rejects_nonfinite_statistics(self):
        from projection_context import validate_context
        codes = {'03-23-N'}
        data = {'day':'2026-06-30','minute':1435,'roomCode':'03-23-N','roomStats':{'mean':24.5},'estimated':['03-23-N']}
        self.assertEqual(validate_context(data,codes)['roomStats']['mean'],24.5)
        for changed in ({'minute':1440},{'roomStats':{'mean':float('nan')}},{'estimated':['unknown']}):
            with self.assertRaises(ValueError):
                validate_context({**data,**changed},codes)

    def test_real_workbook_and_room_codes(self):
        rows = list(read_rows(RESOURCE_DIR / 'température coupé.xlsx'))
        docs = [doc for row in rows if (doc := measurement(row))]
        self.assertEqual(len(rows), 3055)
        self.assertEqual(len(docs), 2944)
        self.assertTrue(any(doc['room'] == '03-23-N' for doc in docs))
        self.assertEqual(docs[0]['temperature'], 16.08)
        self.assertEqual(docs[0]['timestamp'], dt.datetime(2025, 12, 31, 23, 0, 0, 365000, tzinfo=dt.timezone.utc))

    def test_invalid_values_and_truncated_rows(self):
        row = {'timestamp': '2026-01-01 00:00:00.365', 'value': '16.08',
               'room': '00-02-S', 'floor': 'RDC', 'zone': 'ESEO'}
        for value in ('', None, 'nan', 'inf', 'abc'):
            self.assertIsNone(measurement({**row, 'value': value}))
        self.assertIsNone(measurement({**row, 'timestamp': 'invalid'}))
        self.assertIsNone(measurement({'timestamp': row['timestamp']}))

    def test_regular_csv_and_xlsx(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            csv = Path(directory) / 'data.csv'
            csv.write_text('timestamp,value,room,floor,zone\n2026-01-01T00:00:00Z,16.08,03-23-N,N3,ESEO\n')
            self.assertEqual(measurement(next(read_rows(csv)))['temperature'], 16.08)
            workbook = Workbook()
            workbook.active.append(['timestamp', 'value', 'room', 'floor', 'zone'])
            workbook.active.append([dt.datetime(2026, 1, 1), 16.08, '03-23-N', 'N3', 'ESEO'])
            xlsx = Path(directory) / 'data.xlsx'
            workbook.save(xlsx)
            self.assertEqual(measurement(next(read_rows(xlsx)))['room'], '03-23-N')

    def test_embedded_csv_with_extra_empty_excel_columns(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            workbook = Workbook()
            workbook.active.append(['timestamp,value,room,floor,zone'])
            workbook.active.append(['2026-01-01T00:00:00Z,16.08,03-23-N,N3,ESEO'])
            # Une cellule formatée suffit à gonfler les dimensions Excel.
            workbook.active.cell(5, 3).number_format = '0.00'
            path = Path(directory) / 'full.xlsx'
            workbook.save(path)
            rows = list(read_rows(path))
            self.assertEqual(len(rows), 1)
            self.assertEqual(measurement(rows[0])['temperature'], 16.08)


class ApiTests(unittest.TestCase):
    def test_projection_context_forwarded_with_heatmap(self):
        payload = {'temperatures':{'03-23-N':24.5}, 'context':{'day':'2026-06-30','minute':1435,
                   'estimated':['03-23-N'],'roomCode':'03-23-N','roomStats':{'min':21,'mean':23,'max':25}}}
        with patch.object(self.module.requests,'post') as post:
            response = self.client.post('/api/project-temperatures',json=payload,headers=self.headers)
            self.assertEqual(response.status_code,200)
            self.assertEqual(post.call_args.kwargs['json'],payload)
        payload['context']['minute']=1440
        self.assertEqual(self.client.post('/api/project-temperatures',json=payload,headers=self.headers).status_code,400)

    @classmethod
    def setUpClass(cls):
        with patch.dict(os.environ, {'JWT_SECRET': 'test-secret'}), patch('pymongo.MongoClient'):
            cls.module = importlib.import_module('app')
        cls.client = cls.module.app.test_client()
        cls.headers = {'Authorization': 'Bearer ' + cls.module.make_token('test')}

    def test_auth_and_timestamp_validation(self):
        self.assertEqual(self.client.get('/api/temperatures').status_code, 401)
        for timestamp in ('', 'bad', '2026-01-01T00:00:00'):
            response = self.client.get('/api/temperatures', query_string={'timestamp': timestamp}, headers=self.headers)
            self.assertEqual(response.status_code, 400)

    def test_daily_views_require_existing_authentication(self):
        for path in ('/api/plans', '/api/days/3', '/api/day/3/2026-01-01'):
            self.assertEqual(self.client.get(path).status_code, 401)

    def test_range_empty_and_populated(self):
        collection = MagicMock()
        with patch.object(self.module, 'temperatures', collection):
            collection.find_one.return_value = None
            self.assertEqual(self.client.get('/api/temperatures/range', headers=self.headers).json,
                             {'minTimestamp': None, 'maxTimestamp': None})
            collection.find_one.return_value = {'timestamp': dt.datetime(2026, 1, 1)}
            self.assertEqual(self.client.get('/api/temperatures/range', headers=self.headers).json['minTimestamp'],
                             '2026-01-01T00:00:00+00:00')

    def test_temperature_response_and_window(self):
        collection = MagicMock()
        collection.aggregate.return_value = [{'_id': '03-23-N', 'temperature': 16.08}]
        with patch.object(self.module, 'temperatures', collection):
            response = self.client.get('/api/temperatures', query_string={'timestamp': '2026-01-01T00:00:00Z'}, headers=self.headers)
        self.assertEqual(response.json, {'03-23-N': 16.08})
        window = collection.aggregate.call_args.args[0][0]['$match']['timestamp']
        self.assertEqual(window['$lte'] - window['$gte'], dt.timedelta(minutes=20))


    def test_project_temperatures_forwards_and_clears(self):
        for values in ({'03-23-N': 21}, {}):
            with patch.object(self.module.requests, 'post') as post:
                response = self.client.post('/api/project-temperatures', json=values, headers=self.headers)
            self.assertEqual(response.status_code, 200)
            post.assert_called_once_with(f'{self.module.CAMERA_URL}/project-temperatures', json=values, timeout=5)

    def test_project_temperatures_validation_and_auth(self):
        self.assertEqual(self.client.post('/api/project-temperatures', json={}).status_code, 401)
        with patch.object(self.module.requests, 'post') as post:
            for values in ([], {'unknown': 21}, {'03-23-N': True}, {'03-23-N': '21'}, {'03-23-N': float('inf')}):
                response = self.client.post('/api/project-temperatures', json=values, headers=self.headers)
                self.assertEqual(response.status_code, 400)
            post.assert_not_called()

    def test_unavailable_projector(self):
        with patch.object(self.module.requests, 'post', side_effect=self.module.requests.ConnectionError):
            response = self.client.post('/api/project-temperatures', json={}, headers=self.headers)
        self.assertEqual(response.status_code, 502)

    def test_calibration_settings_auth_and_forwarding(self):
        self.assertEqual(self.client.get('/api/calibration-settings').status_code, 401)
        self.assertEqual(self.client.post('/api/calibration-settings', json={}).status_code, 401)
        settings = {'mode': 'manual', 'corners': [[0, 0], [1, 0], [1, 1], [0, 1]], 'showGuide': True}
        with patch.object(self.module.requests, 'post') as post:
            post.return_value.status_code = 200
            post.return_value.json.return_value = {**settings, 'automaticReady': False}
            response = self.client.post('/api/calibration-settings', json=settings, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['mode'], 'manual')
        post.assert_called_once_with(f'{self.module.CAMERA_URL}/calibration-settings', json=settings, timeout=5)
        with patch.object(self.module.requests, 'get') as get:
            get.return_value.status_code = 200
            get.return_value.json.return_value = settings
            self.assertEqual(self.client.get('/api/calibration-settings', headers=self.headers).json, settings)

    def test_projection_preview_auth_and_image_forwarding(self):
        self.assertEqual(self.client.get('/api/projection-preview').status_code, 401)
        with patch.object(self.module.requests, 'get') as get:
            get.return_value.content = b'jpeg-test'
            get.return_value.headers = {'X-Calibration-Settings': '{"settingsRevision":4}'}
            response = self.client.get('/api/projection-preview', headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data, b'jpeg-test')
        self.assertEqual(response.mimetype, 'image/jpeg')
        self.assertEqual(response.headers['Cache-Control'], 'no-store')
        self.assertEqual(response.headers['X-Calibration-Settings'], '{"settingsRevision":4}')
        with patch.object(self.module.requests, 'get', side_effect=self.module.requests.ConnectionError):
            self.assertEqual(self.client.get('/api/projection-preview', headers=self.headers).status_code, 502)

    def test_invalid_calibration_and_unavailable_receiver(self):
        with patch.object(self.module.requests, 'post') as post:
            self.assertEqual(self.client.post('/api/calibration-settings', json=[], headers=self.headers).status_code, 400)
            post.assert_not_called()
            post.return_value.status_code = 400
            post.return_value.json.return_value = {'error': 'Coins croisés.'}
            response = self.client.post('/api/calibration-settings', json={}, headers=self.headers)
            self.assertEqual(response.status_code, 400)
            self.assertEqual(response.json['error'], 'Coins croisés.')
        with patch.object(self.module.requests, 'get', side_effect=self.module.requests.ConnectionError):
            self.assertEqual(self.client.get('/api/calibration-settings', headers=self.headers).status_code, 502)


if __name__ == '__main__':
    unittest.main()
